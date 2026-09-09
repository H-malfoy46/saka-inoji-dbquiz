"""
楽曲DB・メンバーDBのExcelを読んで、クイズ用データ(quiz_data.js)を生成する。

Excelを正にしているので、DBを更新したらこのスクリプトを再実行するだけで
クイズの中身も自動的に最新になる。
"""
import json
import os
import re
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.abspath(__file__))
OTHER = os.path.dirname(BASE)

SONG_XLSX = os.path.join(OTHER, "坂道イコノイジョイ楽曲DB", "エクセルデータ(メイン)",
                         "坂道イコノイジョイ楽曲データベース.xlsx")
MEMBER_XLSX = os.path.join(OTHER, "坂道イコノイジョイメンバーDB", "エクセルデータ(メイン)",
                           "坂道イコノイジョイメンバーデータベース.xlsx")

SONG_KEYS = ["title", "cd", "date", "type", "center", "lyricist", "composer",
             "arranger", "choreographer", "director", "members", "notes", "sources"]
MEMBER_KEYS = ["name", "period", "status", "tenure", "birthday", "color",
               "center_songs", "notes"]

# クイズに使ってはいけない値(不明・未確定・プレースホルダ)
BAD_PAT = re.compile(r"要確認|不明|未確定|該当なし|^[-‐―–—ー]$|^$")


def is_usable(v):
    """クイズの答え・選択肢として使える値か"""
    if v is None:
        return False
    s = str(v).strip()
    if not s or BAD_PAT.search(s):
        return False
    if s.startswith("(") or s.startswith("（"):  # 「(14名)」のような注記
        return False
    return True


def is_single_person(v):
    """答えが一意になるよう、単独名義のものだけ採用する"""
    if not is_usable(v):
        return False
    s = str(v)
    # 複数人が併記されている区切り記号
    for sep in ["、", "，", ",", "・", "／", "/", "&", "＆", " と ", "＋", "+"]:
        if sep in s:
            return False
    if len(s) > 24:
        return False
    return True


def read_sheet(path, keys, ncols):
    """各シートを読む。本表の下にある『参考』表は空行で切れるので、そこで読み終える。"""
    wb = load_workbook(path, data_only=True)
    out = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = []
        for r in range(3, ws.max_row + 1):
            first = ws.cell(row=r, column=2).value
            if first is None or str(first).strip() == "":
                break  # 本表おわり(以降は参考表)
            rec = {}
            for i, k in enumerate(keys):
                v = ws.cell(row=r, column=2 + i).value
                rec[k] = "" if v is None else str(v).strip()
            rows.append(rec)
        out[name] = rows
    return out


ORD_RE = re.compile(r"\((\d+)(?:st|nd|rd|th)(アルバム|ミニアルバム)?\)$")


def parse_cd(cd):
    """CD名から「何thシングル/アルバムか」を取り出す"""
    m = ORD_RE.search(cd)
    if not m:
        return None, None
    n = int(m.group(1))
    kind = m.group(2) or "シングル"
    return n, kind


def load_yomi():
    """メンバーの「よみ」はExcelに列が無いので、元のJSONから拾って五十音順の並べ替えに使う"""
    import glob
    yomi = {}
    pat = os.path.join(OTHER, "坂道イコノイジョイメンバーDB", "関連資料", "member_*.json")
    for path in glob.glob(pat):
        try:
            with open(path, "r", encoding="utf-8") as f:
                for r in json.load(f):
                    if r.get("name") and r.get("yomi"):
                        yomi[r["name"]] = r["yomi"]
        except Exception:
            pass
    return yomi


def main():
    songs_by_group = read_sheet(SONG_XLSX, SONG_KEYS, len(SONG_KEYS))
    members_by_group = read_sheet(MEMBER_XLSX, MEMBER_KEYS, len(MEMBER_KEYS))
    yomi_map = load_yomi()

    songs = []
    for g, rows in songs_by_group.items():
        for i, r in enumerate(rows):
            n, kind = parse_cd(r["cd"])
            songs.append({
                "g": g,
                "idx": i,          # Excel上の並び順(発売日昇順)
                "title": r["title"],
                "cd": r["cd"],
                "cdBase": r["cd"].split("(")[0].strip(),
                "num": n,
                "kind": kind,
                "date": r["date"],
                "year": (r["date"][:4] if re.match(r"^\d{4}", r["date"] or "") else None),
                "type": r["type"],
                "center": r["center"],
                "lyricist": r["lyricist"],
                "composer": r["composer"],
                "arranger": r["arranger"],
                "choreographer": r["choreographer"],
                "director": r["director"],
                "members": r["members"],
                "notes": r["notes"],
                "sources": r["sources"],
            })

    members = []
    name_groups = {}
    for g, rows in members_by_group.items():
        for i, r in enumerate(rows):
            members.append({
                "g": g,
                "idx": i,          # Excel上の並び順(期→五十音)
                "name": r["name"],
                "yomi": yomi_map.get(r["name"], ""),
                "period": r["period"],
                "status": r["status"],
                "tenure": r["tenure"],
                "birthday": r["birthday"],
                "color": r["color"],
                "centerSongs": r["center_songs"],
                "notes": r["notes"],
            })
            name_groups.setdefault(r["name"], set()).add(g)

    # 複数グループに在籍した人(欅坂46→櫻坂46 など)は「所属グループ当て」に使えない
    for m in members:
        m["uniqueGroup"] = len(name_groups[m["name"]]) == 1

    data = {
        "groups": list(songs_by_group.keys()),
        "songs": songs,
        "members": members,
    }

    out_path = os.path.join(BASE, "quiz_data.js")
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("window.QUIZ_DATA = ")
        json.dump(data, f, ensure_ascii=False, separators=(",", ":"))
        f.write(";\n")

    # ---- 生成結果のかんたんな検算 ----
    print("saved:", out_path)
    print(f"  楽曲 {len(songs)}件 / メンバー {len(members)}件")
    print()
    fields = [("center", "センター"), ("composer", "作曲"), ("arranger", "編曲"),
              ("lyricist", "作詞"), ("choreographer", "振付"), ("director", "MV監督")]
    print("  [楽曲] 単独名義で出題に使える件数")
    for k, label in fields:
        c = sum(1 for s in songs if is_single_person(s[k]))
        print(f"    {label}: {c}")
    print(f"    何thシングル判定可: {sum(1 for s in songs if s['num'])}")
    print()
    print("  [メンバー] 出題に使える件数")
    for k, label in [("period", "期"), ("color", "カラー"), ("birthday", "誕生日"),
                     ("centerSongs", "センター曲")]:
        c = sum(1 for m in members if is_usable(m[k]))
        print(f"    {label}: {c}")
    print(f"    所属1グループのみ(グループ当て可): {sum(1 for m in members if m['uniqueGroup'])}")


if __name__ == "__main__":
    main()
